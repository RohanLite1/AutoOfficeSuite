"""
Estimate Parts Handler Module.
Extracts parts line items & amounts from bill estimate PDFs
and writes them directly into the PARTS COMPUTATION table of Master Assessment Excels (.xlsx, .xlsm, .xls).
"""

import os
import re
import shutil
from datetime import datetime
import pdfplumber

# In-memory cache for inspect_estimate_excel: (norm_path, sheet_name) -> (mtime, size, info_dict)
_ESTIMATE_INFO_CACHE = {}


def _prune_old_backups(file_path, max_keep=5):
    """Keep only the latest N backups of the specified file to prevent disk bloat."""
    try:
        dir_name = os.path.dirname(file_path)
        base_name = os.path.basename(file_path)
        prefix = f"{base_name}.bak_"
        backups = []
        for fname in os.listdir(dir_name):
            if fname.startswith(prefix):
                full_p = os.path.join(dir_name, fname)
                try:
                    backups.append((os.path.getmtime(full_p), full_p))
                except OSError:
                    pass
        backups.sort(key=lambda x: x[0], reverse=True)
        for _, old_backup in backups[max_keep:]:
            try:
                os.remove(old_backup)
            except OSError:
                pass
    except Exception:
        pass


def invalidate_estimate_cache(file_path=None):
    """Invalidate cache for a specific file or all files."""
    global _ESTIMATE_INFO_CACHE
    if file_path:
        norm = os.path.normpath(os.path.abspath(file_path)).lower()
        keys_to_del = [k for k in list(_ESTIMATE_INFO_CACHE.keys()) if k[0].lower() == norm]
        for k in keys_to_del:
            _ESTIMATE_INFO_CACHE.pop(k, None)
    else:
        _ESTIMATE_INFO_CACHE.clear()


def parse_estimate_pdf(pdf_path):
    """
    Parses a dealer bill estimate PDF (e.g. BMW, Maruti, Hyundai, etc.)
    Returns header info and list of extracted line items (Parts and Labour).
    """
    if not os.path.exists(pdf_path):
        raise FileNotFoundError(f"PDF not found: {pdf_path}")

    header_info = {
        "invoice_no": "",
        "invoice_date": "",
        "registration_no": "",
        "make": "",
        "model": "",
        "service_advisor": "",
        "dealer_name": ""
    }
    items = []

    with pdfplumber.open(pdf_path) as pdf:
        for p_idx, page in enumerate(pdf.pages):
            text = page.extract_text() or ""
            lines = text.splitlines()

            # Header details on first page
            if p_idx == 0:
                # Reg no (Indian registration number format)
                m_reg = re.search(r'\b([A-Z]{2}\s*\d{1,2}\s*[A-Z]{1,3}\s*\d{4})\b', text)
                if m_reg: header_info['registration_no'] = m_reg.group(1).replace(' ', '')

                # Invoice no
                m_inv = re.search(r'\b([A-Z]{3,5}\d{8,12})\b', text)
                if m_inv: header_info['invoice_no'] = m_inv.group(1)

                # Date
                m_date = re.search(r'Invoice\s*Date\s*([0-9]{1,2}-[A-Za-z]{3}-[0-9]{2,4})', text, re.IGNORECASE)
                if m_date: header_info['invoice_date'] = m_date.group(1)

                # Advisor
                m_adv = re.search(r'Service\s*Advisor\s*([A-Za-z\s]+?)(?:,|$|\n)', text, re.IGNORECASE)
                if m_adv: header_info['service_advisor'] = m_adv.group(1).strip()

                # Dealer name
                for line in lines[:5]:
                    line_s = line.strip()
                    if any(term in line_s.lower() for term in ['automobiles', 'motors', 'bmw', 'maruti', 'hyundai', 'workshop', 'services']):
                        header_info['dealer_name'] = line_s.split('   ')[0].strip()
                        break

            # Parse line items across pages
            for line in lines:
                line_clean = line.strip()

                # Pattern 1: Krishna Automobiles / Dealer Line:
                # No. (code) | Description | HSN | Qty | UoM | Unit Price | Disc % | Discount | CGST | SGST | Amount
                # e.g.: 83192289285 Repair flat 35069190 2 PCS 6775.00 10.0% 1355.00 9.0% 9.0% 12195.00
                m1 = re.match(
                    r"^([A-Z0-9\-\/]{4,15})\s+(.+?)\s+(\d{4,8})\s+(\d+(?:\.\d+)?)\s+([A-Za-z]+)\s+([\d,]+\.\d{2})\s+.*?\s+([\d,]+\.\d{2})$",
                    line_clean
                )
                if m1:
                    code, desc, hsn, qty, uom, unit_price, amt = m1.groups()
                    desc_clean = desc.strip()
                    is_labour = "FRU" in uom.upper() or "LABOR" in desc_clean.upper() or "JOB" in desc_clean.upper() or "R / BODY" in desc_clean.upper()
                    items.append({
                        "code": code.strip(),
                        "particular": desc_clean,
                        "hsn": hsn.strip(),
                        "qty": f"{qty} {uom}",
                        "unit_price": float(unit_price.replace(",", "")),
                        "amount": float(amt.replace(",", "")),
                        "item_type": "Labour" if is_labour else "Part",
                        "selected": not is_labour  # Default select parts
                    })
                    continue

                # Pattern 2: Standard Line Item with Part Code, Description, Quantity & Amount
                # e.g. 51319478772 Frontscheibe 1 PCS 48086.32
                m2 = re.match(
                    r"^([A-Z0-9\-\/]{4,15})\s+([A-Za-z0-9\s\.\-\/\(\)&]{3,40}?)\s+(\d+(?:\.\d+)?)\s*([A-Za-z]{2,5})?\s+([\d,]+\.\d{2})$",
                    line_clean
                )
                if m2:
                    code, desc, qty, uom, amt = m2.groups()
                    desc_clean = desc.strip()
                    if desc_clean.lower() not in ["total", "cgst", "sgst", "igst", "sub total", "amount to pay", "round off", "sum parts"]:
                        if not any(k in desc_clean.lower() for k in ["terms", "signature", "authorized", "customer"]):
                            is_labour = "LABOR" in desc_clean.upper() or (uom and "FRU" in uom.upper())
                            items.append({
                                "code": code.strip(),
                                "particular": desc_clean,
                                "hsn": "",
                                "qty": f"{qty} {uom or ''}".strip(),
                                "unit_price": float(amt.replace(",", "")),
                                "amount": float(amt.replace(",", "")),
                                "item_type": "Labour" if is_labour else "Part",
                                "selected": not is_labour
                            })

    return {
        "success": True,
        "filename": os.path.basename(pdf_path),
        "header": header_info,
        "items": items,
        "total_parts": sum(1 for it in items if it["item_type"] == "Part"),
        "total_estimate_sum": sum(it["amount"] for it in items if it["item_type"] == "Part")
    }


def inspect_estimate_excel(file_path, target_sheet=None):
    """
    Inspects an Excel spreadsheet (.xlsx, .xlsm, or .xls) to locate the active PARTS COMPUTATION table.
    Scans VISIBLE rows to locate the visible table header and first visible data row (e.g. Row 171).
    Uses mtime caching for instant responses.
    """
    if not os.path.exists(file_path):
        return {"success": False, "error": f"File does not exist: {file_path}"}

    norm_path = os.path.normpath(os.path.abspath(file_path))
    try:
        stat = os.stat(norm_path)
        mtime, size = stat.st_mtime, stat.st_size
    except OSError:
        return {"success": False, "error": f"Cannot access file: {file_path}"}

    cache_key = (norm_path, str(target_sheet or "").strip())
    if cache_key in _ESTIMATE_INFO_CACHE:
        cached_mtime, cached_size, cached_info = _ESTIMATE_INFO_CACHE[cache_key]
        if cached_mtime == mtime and cached_size == size:
            return cached_info.copy()

    ext = os.path.splitext(norm_path)[1].lower()

    if ext == ".xls":
        try:
            import xlrd
            wb = xlrd.open_workbook(file_path, formatting_info=True)
            sheet_names = wb.sheet_names()

            best_sheet = ""
            best_score = -1
            best_header_r = 169
            best_start_r = 171
            best_cols = {"es": 1, "bill_no": 2, "particular": 3, "estimate": 4}

            for sname in sheet_names:
                s = wb.sheet_by_name(sname)
                # Ignore hidden sheets like 'data' unless explicitly requested
                if s.visibility != 0 and sname != target_sheet:
                    continue

                score = 0
                h_row = -1
                cols = {"es": 1, "bill_no": 2, "particular": 3, "estimate": 4}

                # Scan through all rows in sheet, skipping hidden rows
                for r in range(s.nrows):
                    rinfo = s.rowinfo_map.get(r) if hasattr(s, "rowinfo_map") else None
                    if rinfo and rinfo.hidden:
                        continue  # Skip hidden rows!

                    row_vals = [str(s.cell_value(r, c)).strip() for c in range(min(20, s.ncols))]
                    row_str = " ".join(row_vals).upper()

                    if "PARTS COMPUTATION" in row_str:
                        score += 50
                        # Look at subsequent visible rows for Particular / Estimate
                        for r2 in range(r + 1, min(r + 15, s.nrows)):
                            r2info = s.rowinfo_map.get(r2) if hasattr(s, "rowinfo_map") else None
                            if r2info and r2info.hidden:
                                continue
                            r2_vals = [str(s.cell_value(r2, c)).strip() for c in range(min(20, s.ncols))]
                            if any("PARTICULAR" in v.upper() for v in r2_vals):
                                h_row = r2 + 1
                                for c, v in enumerate(r2_vals):
                                    v_u = v.upper()
                                    if "PARTICULAR" in v_u: cols["particular"] = c + 1
                                    elif "ESTIMATE" in v_u: cols["estimate"] = c + 1
                                    elif "BILL" in v_u: cols["bill_no"] = c + 1
                                    elif "ES." in v_u or v_u == "ES" or "S.NO" in v_u: cols["es"] = c + 1
                                break
                        if h_row > 0:
                            break

                    elif any("PARTICULAR" in v.upper() for v in row_vals) and any("ESTIMATE" in v.upper() for v in row_vals):
                        score += 40
                        h_row = r + 1
                        for c, v in enumerate(row_vals):
                            v_u = v.upper()
                            if "PARTICULAR" in v_u: cols["particular"] = c + 1
                            elif "ESTIMATE" in v_u: cols["estimate"] = c + 1
                            elif "BILL" in v_u: cols["bill_no"] = c + 1
                            elif "ES." in v_u or v_u == "ES" or "S.NO" in v_u: cols["es"] = c + 1
                        break

                if sname.lower() == "sheet1":
                    score += 15
                if target_sheet and sname.strip().lower() == target_sheet.strip().lower():
                    score += 200

                if score > best_score:
                    best_score = score
                    best_sheet = sname
                    best_cols = cols
                    if h_row > 0:
                        best_header_r = h_row
                        # Find the first VISIBLE data row following the header row
                        best_start_r = h_row + 1
                        for r_data in range(h_row, min(h_row + 15, s.nrows)):
                            r_d_info = s.rowinfo_map.get(r_data) if hasattr(s, "rowinfo_map") else None
                            if r_d_info and r_d_info.hidden:
                                continue
                            sub_vals = [str(s.cell_value(r_data, c)).strip().lower() for c in range(min(15, s.ncols))]
                            if any(term in sub_vals for term in ["metal", "non metal", "glass", "labour", "tax"]):
                                continue
                            best_start_r = r_data + 1
                            break

            if not best_sheet and sheet_names:
                best_sheet = sheet_names[0]

            res = {
                "success": True,
                "file_path": norm_path,
                "file_type": "xls",
                "sheets": sheet_names,
                "selected_sheet": best_sheet,
                "header_row": best_header_r,
                "data_start_row": best_start_r,
                "particular_col": best_cols["particular"],
                "estimate_col": best_cols["estimate"],
                "bill_no_col": best_cols["bill_no"],
                "es_col": best_cols["es"],
                "particular_col_letter": chr(64 + best_cols["particular"]) if 1 <= best_cols["particular"] <= 26 else "C",
                "estimate_col_letter": chr(64 + best_cols["estimate"]) if 1 <= best_cols["estimate"] <= 26 else "D"
            }
            _ESTIMATE_INFO_CACHE[cache_key] = (mtime, size, res)
            return res.copy()
        except Exception as e:
            return {"success": False, "error": f"Failed reading .xls file: {str(e)}"}

    else:
        # .xlsx or .xlsm
        wb = None
        try:
            import openpyxl
            wb = openpyxl.load_workbook(norm_path, data_only=True)
            sheet_names = wb.sheetnames

            best_sheet = ""
            best_score = -1
            best_header_r = 169
            best_start_r = 171
            best_cols = {"es": 1, "bill_no": 2, "particular": 3, "estimate": 4}

            for sname in sheet_names:
                sheet = wb[sname]
                score = 0
                h_row = -1
                cols = {"es": 1, "bill_no": 2, "particular": 3, "estimate": 4}

                for r in range(1, min(sheet.max_row + 1, 600)):
                    dim = sheet.row_dimensions.get(r)
                    if dim and dim.hidden:
                        continue  # Skip hidden row!

                    row_vals = [str(sheet.cell(row=r, column=c).value or "").strip() for c in range(1, min(20, sheet.max_column + 1))]
                    row_str = " ".join(row_vals).upper()

                    if "PARTS COMPUTATION" in row_str:
                        score += 50
                        for r2 in range(r + 1, min(r + 15, sheet.max_row + 1)):
                            dim2 = sheet.row_dimensions.get(r2)
                            if dim2 and dim2.hidden:
                                continue
                            r2_vals = [str(sheet.cell(row=r2, column=c).value or "").strip() for c in range(1, min(20, sheet.max_column + 1))]
                            if any("PARTICULAR" in v.upper() for v in r2_vals):
                                h_row = r2
                                for c, v in enumerate(r2_vals):
                                    v_u = v.upper()
                                    if "PARTICULAR" in v_u: cols["particular"] = c + 1
                                    elif "ESTIMATE" in v_u: cols["estimate"] = c + 1
                                    elif "BILL" in v_u: cols["bill_no"] = c + 1
                                    elif "ES." in v_u or v_u == "ES" or "S.NO" in v_u: cols["es"] = c + 1
                                break
                        if h_row > 0:
                            break

                    elif any("PARTICULAR" in v.upper() for v in row_vals) and any("ESTIMATE" in v.upper() for v in row_vals):
                        score += 40
                        h_row = r
                        for c, v in enumerate(row_vals):
                            v_u = v.upper()
                            if "PARTICULAR" in v_u: cols["particular"] = c + 1
                            elif "ESTIMATE" in v_u: cols["estimate"] = c + 1
                            elif "BILL" in v_u: cols["bill_no"] = c + 1
                            elif "ES." in v_u or v_u == "ES" or "S.NO" in v_u: cols["es"] = c + 1
                        break

                if sname.lower() == "sheet1":
                    score += 15
                if target_sheet and sname.strip().lower() == target_sheet.strip().lower():
                    score += 200

                if score > best_score:
                    best_score = score
                    best_sheet = sname
                    best_cols = cols
                    if h_row > 0:
                        best_header_r = h_row
                        best_start_r = h_row + 1
                        for r_data in range(h_row + 1, min(h_row + 15, sheet.max_row + 1)):
                            dim_d = sheet.row_dimensions.get(r_data)
                            if dim_d and dim_d.hidden:
                                continue
                            sub_vals = [str(sheet.cell(row=r_data, column=c).value or "").strip().lower() for c in range(1, min(15, sheet.max_column + 1))]
                            if any(term in sub_vals for term in ["metal", "non metal", "glass", "labour", "tax"]):
                                continue
                            best_start_r = r_data
                            break

            if not best_sheet and sheet_names:
                best_sheet = sheet_names[0]

            res = {
                "success": True,
                "file_path": norm_path,
                "file_type": "xlsx",
                "sheets": sheet_names,
                "selected_sheet": best_sheet,
                "header_row": best_header_r,
                "data_start_row": best_start_r,
                "particular_col": best_cols["particular"],
                "estimate_col": best_cols["estimate"],
                "bill_no_col": best_cols["bill_no"],
                "es_col": best_cols["es"],
                "particular_col_letter": openpyxl.utils.get_column_letter(best_cols["particular"]),
                "estimate_col_letter": openpyxl.utils.get_column_letter(best_cols["estimate"])
            }
            _ESTIMATE_INFO_CACHE[cache_key] = (mtime, size, res)
            return res.copy()
        except Exception as e:
            return {"success": False, "error": f"Failed reading Excel file: {str(e)}"}
        finally:
            if wb is not None:
                try:
                    wb.close()
                except Exception:
                    pass


def write_parts_to_excel(file_path, sheet_name=None, parts=None, start_row=None, bill_no=""):
    """
    Writes extracted parts into the PARTS COMPUTATION table:
    - Serial No ('Es.'): Sequential item index (e.g. 1, 2, 3...)
    - Bill No ('Bill No.'): Optional bill reference
    - Particular ('Particular'): Part description (Col C)
    - Estimate ('Estimate'): Part Amount (Col D)
    Ensures parts are written strictly to VISIBLE rows, skipping any hidden template rows.
    """
    if not os.path.exists(file_path):
        return {"success": False, "error": f"File not found: {file_path}"}

    if not parts:
        return {"success": False, "error": "No parts provided to write."}

    # Safety check for file lock
    try:
        with open(file_path, "r+b"):
            pass
    except PermissionError:
        return {
            "success": False,
            "error": "The Excel file is currently open in Microsoft Excel. Please close it first."
        }
    except Exception as e:
        return {"success": False, "error": f"File access error: {str(e)}"}

    # Create safety backup
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_path = f"{file_path}.bak_{timestamp}"
    try:
        shutil.copy2(file_path, backup_path)
        _prune_old_backups(file_path)
    except Exception as e:
        print(f"Warning creating backup: {e}")

    # Inspect sheet and determine exact target columns and start row
    sheet_info = inspect_estimate_excel(file_path, target_sheet=sheet_name)
    if not sheet_info.get("success"):
        return {"success": False, "error": sheet_info.get("error", "Failed to inspect Excel file.")}

    target_sheet_name = sheet_info.get("selected_sheet")
    part_col_1based = sheet_info.get("particular_col", 3)
    est_col_1based = sheet_info.get("estimate_col", 4)
    bill_col_1based = sheet_info.get("bill_no_col", 2)
    es_col_1based = sheet_info.get("es_col", 1)

    # Determine start row (1-indexed). If not specified or if set to hidden rows (like 7 or 28), use detected visible start row!
    detected_visible_start = sheet_info.get("data_start_row", 171)
    if start_row is None or int(start_row) <= 0:
        start_row = detected_visible_start
    else:
        start_row = int(start_row)
        # If user had an old start_row like 7 or 28 from earlier runs that is hidden, snap to detected visible start row
        if start_row < detected_visible_start and detected_visible_start >= 100:
            start_row = detected_visible_start

    ext = os.path.splitext(file_path)[1].lower()

    if ext == ".xls":
        # Legacy .xls using xlrd + xlutils.copy
        try:
            import xlrd
            from xlutils.copy import copy

            rb = xlrd.open_workbook(file_path, formatting_info=True)
            sheet_names = rb.sheet_names()

            sheet_idx = 0
            if target_sheet_name in sheet_names:
                sheet_idx = sheet_names.index(target_sheet_name)
            else:
                for idx, sn in enumerate(sheet_names):
                    if sn.strip().lower() == str(target_sheet_name).strip().lower():
                        sheet_idx = idx
                        target_sheet_name = sn
                        break

            rb_sheet = rb.sheet_by_index(sheet_idx)
            wb = copy(rb)
            sheet = wb.get_sheet(sheet_idx)

            # 0-indexed rows and columns in xlwt
            current_r = max(0, start_row - 1)
            c_part = max(0, part_col_1based - 1)
            c_est = max(0, est_col_1based - 1)
            c_bill = max(0, bill_col_1based - 1)
            c_es = max(0, es_col_1based - 1)

            item_num = 1
            written = 0
            first_written_row = None
            last_written_row = None

            for p in parts:
                desc = p.get("particular", "").strip()
                amt = p.get("amount", 0.0)
                try:
                    amt_num = float(amt)
                except (ValueError, TypeError):
                    amt_num = 0.0

                # Ensure current_r is a VISIBLE row!
                while current_r < rb_sheet.nrows:
                    rinfo = rb_sheet.rowinfo_map.get(current_r) if hasattr(rb_sheet, "rowinfo_map") else None
                    if rinfo and rinfo.hidden:
                        current_r += 1
                    else:
                        break

                if first_written_row is None:
                    first_written_row = current_r + 1
                last_written_row = current_r + 1

                # Write item serial number in Col Es.
                try:
                    sheet.write(current_r, c_es, item_num)
                except Exception:
                    pass

                # Write Bill No if provided
                if bill_no:
                    sheet.write(current_r, c_bill, str(bill_no).strip())

                # Write Particular (Part Description) into Col C
                sheet.write(current_r, c_part, desc)

                # Write Estimate Amount into Col D
                sheet.write(current_r, c_est, amt_num)

                current_r += 1
                item_num += 1
                written += 1

            wb.save(file_path)
            invalidate_estimate_cache(file_path)

            return {
                "success": True,
                "written_count": written,
                "start_row": first_written_row or start_row,
                "end_row": last_written_row or (start_row + written - 1),
                "backup_created": backup_path,
                "file_path": file_path,
                "sheet": target_sheet_name,
                "particular_col": chr(65 + c_part),
                "estimate_col": chr(65 + c_est)
            }
        except Exception as e:
            return {"success": False, "error": f"Error writing to .xls file: {str(e)}"}

    else:
        # Modern .xlsx / .xlsm using openpyxl
        wb = None
        try:
            import openpyxl

            wb = openpyxl.load_workbook(file_path)
            if target_sheet_name in wb.sheetnames:
                sheet = wb[target_sheet_name]
            else:
                matched = False
                for sn in wb.sheetnames:
                    if sn.strip().lower() == str(target_sheet_name).strip().lower():
                        sheet = wb[sn]
                        target_sheet_name = sn
                        matched = True
                        break
                if not matched:
                    sheet = wb.active
                    target_sheet_name = sheet.title

            current_r = start_row
            c_part = part_col_1based
            c_est = est_col_1based
            c_bill = bill_col_1based
            c_es = es_col_1based

            item_num = 1
            written = 0
            first_written_row = None
            last_written_row = None

            for p in parts:
                desc = p.get("particular", "").strip()
                amt = p.get("amount", 0.0)
                try:
                    amt_num = float(amt)
                except (ValueError, TypeError):
                    amt_num = 0.0

                # Ensure current_r is a VISIBLE row!
                while current_r <= sheet.max_row:
                    dim = sheet.row_dimensions.get(current_r)
                    if dim and dim.hidden:
                        current_r += 1
                    else:
                        break

                if first_written_row is None:
                    first_written_row = current_r
                last_written_row = current_r

                if c_es > 0:
                    sheet.cell(row=current_r, column=c_es, value=item_num)
                if bill_no and c_bill > 0:
                    sheet.cell(row=current_r, column=c_bill, value=str(bill_no).strip())
                sheet.cell(row=current_r, column=c_part, value=desc)
                sheet.cell(row=current_r, column=c_est, value=amt_num)

                current_r += 1
                item_num += 1
                written += 1

            wb.save(file_path)
            invalidate_estimate_cache(file_path)

            return {
                "success": True,
                "written_count": written,
                "start_row": first_written_row or start_row,
                "end_row": last_written_row or (start_row + written - 1),
                "backup_created": backup_path,
                "file_path": file_path,
                "sheet": target_sheet_name,
                "particular_col": openpyxl.utils.get_column_letter(c_part),
                "estimate_col": openpyxl.utils.get_column_letter(c_est)
            }
        except Exception as e:
            return {"success": False, "error": f"Error writing to Excel: {str(e)}"}
        finally:
            if wb is not None:
                try:
                    wb.close()
                except Exception:
                    pass

